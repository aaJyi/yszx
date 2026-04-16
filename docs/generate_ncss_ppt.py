# -*- coding: utf-8 -*-
"""生成 NCSS 算法流程 3 页 PPT（健康场景优化版 + 韩忠明等 2016 核心公式）"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "NCSS_3slides_health_scenario.pptx"


def add_title_slide(prs, title: str, subtitle: str = ""):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(1.15))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.space_before = Pt(8)
    return slide


def add_body(slide, lines: list, top=Inches(1.5)):
    box = slide.shapes.add_textbox(Inches(0.55), top, Inches(9), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, bold, size) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.space_after = Pt(5)
        p.level = 0


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ===== 第 1 页 =====
    s1 = add_title_slide(
        prs,
        "NCSS 算法流程（健康场景）— 相似网络与共同邻居",
        "参考文献：韩忠明等. NCSS: 一种快速有效的复杂网络社团划分算法. 中国科学: 信息科学, 2016, 46(4): 431–444",
    )
    add_body(
        s1,
        [
            ("6.4.1 Step 1  构建相似网络", True, 16),
            (
                "输入：用户标签 / 画像特征；输出：无向图 G=(V,E) 的邻接表 adjacency。"
                " 工程复杂度约 O(n\u00b2\u00b7s)（s 为标签交并运算代价），可用倒排索引、MinHash 压缩候选对。",
                False,
                13,
            ),
            ("6.4.2 Step 2  共同邻居（拓扑相似证据）", True, 16),
            (
                "compute_common_neighbors：为后续中心度与结构强度提供基础；工程上预计算 |C_{i,j}|。",
                False,
                13,
            ),
            ("定义 1（论文式 (1)）  节点 i 与 j 的共同邻居集合", True, 14),
            ("C_{i,j} = Neighbor(i) ∩ Neighbor(j)", False, 16),
            ("常用计数：|C_{i,j}| = |Neighbor(i) ∩ Neighbor(j)|", False, 14),
        ],
    )

    # ===== 第 2 页 =====
    s2 = add_title_slide(
        prs,
        "类 PageRank 节点中心度与种子筛选",
        "Step 3–4：compute_node_centrality → select_seed_nodes",
    )
    add_body(
        s2,
        [
            ("6.4.3 Step 3  节点中心度迭代", True, 16),
            (
                "类 PageRank：邻居贡献 × 共同邻居占比；阻尼系数 c\u2208(0,1)（常取 0.8），"
                "收敛阈值 \u03b5 控制迭代。目标：识别潜在核心节点，避免随机种子。",
                False,
                12,
            ),
            ("定义 2（论文式 (2)）  节点 i 的中心度 PRcen(i)", True, 14),
            (
                "PRcen(i) = c \u00b7 \u03a3_{j\u2208Neighbor(i)} PRcen(j) \u00b7 |C_{i,j}| / "
                "( \u03a3_{k\u2208Neighbor(i)} |C_{i,k}| + 1 ) + (1\u2212c)/|V|",
                False,
                13,
            ),
            ("其中：c 为阻尼系数；|V| 为节点数；分母 “+1” 处理孤立点等退化情形。", False, 11),
            ("6.4.4 Step 4  种子筛选（工程三类约束）", True, 16),
            (
                "① PRcen ≥ 均值；② 度数 deg ≥ 平均度；③ 新种子与已选种子不相邻（尽量覆盖不同社团）。",
                False,
                13,
            ),
        ],
    )

    # ===== 第 3 页 =====
    s3 = add_title_slide(
        prs,
        "社团扩张、迁移、小团过滤与工程优化",
        "Step 5–6：community_expansion → filter_small_communities",
    )
    add_body(
        s3,
        [
            ("6.4.5 Step 5  扩张与节点迁移", True, 15),
            (
                "定义 3（论文式 (3)）  依赖度  D_{a,C} = n_{a,C} / k_a"
                "  （n_{a,C}：a 在社团 C 内邻居数；k_a：a 的度）",
                False,
                12,
            ),
            (
                "定义 4（论文式 (4)）  S_{i,C} = \u03a3_{j\u2208C} [ 2|C_{i,j}|"
                " \u2212 (k_{i,j}^{out}/m)\u00b7s_C^{in} \u2212 (k_{i,j}^{in}/m)\u00b7s_C^{out} ]",
                False,
                11,
            ),
            ("m 为全图度之和；k^{in/out}、s^{in/out}_C 含义同原文。", False, 10),
            (
                "扩张：取 S_{i,C}>0；已属他团时比较 D_{i,C_new} 与 D_{i,C_old}，"
                "若对新团依赖更大且 D_{i,C_new}>0.5（论文算法 2）则迁移。",
                False,
                11,
            ),
            ("6.4.6 Step 6  小社团过滤", True, 15),
            ("过小社团并入最大团或回收（如 |C| < θ|V|，文中示例 θ=5%），提升业务可用性。", False, 12),
            ("工程增强（Heap）：边界候选按 S(i,C) 维护大顶堆，取最优 O(log n)，替代反复全量排序。", True, 12),
            ("论文综合复杂度：T_NCSS = O(kn)（k 为核心节点数，k\u226an）", True, 13),
        ],
    )

    prs.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    main()
