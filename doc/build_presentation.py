# -*- coding: utf-8 -*-
"""生成「医视智行」答辩用 PPT（8 页）。"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor


def _set_run_font(run, name: str = "Microsoft YaHei", size_pt: int = 18, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def add_title_slide(prs, title: str, subtitle: str = ""):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            _set_run_font(r, size_pt=32, bold=True)
    if subtitle and len(slide.placeholders) > 1:
        ph = slide.placeholders[1]
        ph.text = subtitle
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                _set_run_font(r, size_pt=16)


def add_bullet_slide(prs, title: str, bullets: list[str]):
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            _set_run_font(r, size_pt=28, bold=True)
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = line
        p.level = 0
        p.space_after = Pt(6)
        for r in p.runs:
            _set_run_font(r, size_pt=16)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "医视智行——基于 heap 优化 NCSS 的\n多模态智能医疗问诊康复系统",
        "多智能体协同 · 健康档案 · RAG 问答 · 视觉康复 · 社团互助",
    )

    add_bullet_slide(
        prs,
        "需求与项目目标",
        [
            "宏观背景：青年群体亚健康、心理与行为管理难、健康信息碎片化，需可及、可持续的数字健康服务。",
            "用户侧：微信小程序问诊、运动/三餐识别、档案与风险提示、病友交流。",
            "功能需求：中英文医学 RAG、多模态采集、疾病风险与建议、NCSS 社团划分、热点资讯与知识更新。",
            "非功能需求：可解释输出、安全与免责边界、模块化部署、可降级与可运维。",
            "竞赛/交付：完整业务闭环 + 算法落地 + 文档与可复现工程结构。",
        ],
    )

    add_bullet_slide(
        prs,
        "技术栈总览",
        [
            "终端与后台：uni-app 小程序（fronted-mini）、Vue3 管理端（backend-PC）、Spring Boot + MyBatis-Plus（backend）。",
            "AI 智能体：ai-agent（多模态建档、健康分析/建议、与档案 API 对接）。",
            "问答与知识：ai-doctor-rag（向量检索 + 中英文语料、Chroma 等）。",
            "视觉能力：ai-exercise-train（运动识别/计数）、food-train（三餐与营养估计）。",
            "社群与算法：club-division（NCSS 社团划分、相似图与混合边权）。",
            "数据与运维：MySQL、爬虫脚本（热点疾病/头条资讯）、定时与重建任务。",
        ],
    )

    add_bullet_slide(
        prs,
        "ai-agent 智能体工作流程",
        [
            "多模态源数据：体检/病历图片、结构化 raw-health-data 等进入解码与内容构建（文本 + 图像）。",
            "健康档案：LLM 整理为结构化档案 JSON，回填后端；含基础信息、就诊、检验与量表等字段。",
            "健康总结与建议：基于档案 + 原始多模态生成结构化分析 JSON与健康建议，并提交后端。",
            "疾病预测/风险线索：分析结果中的异常指标、维度分析与 extendedData 等作为后续推断与向量特征输入。",
            "编排特点：管道化日志、JSON 修复与字段规范化，保证与 Java DTO 一致、可落库。",
        ],
    )

    add_bullet_slide(
        prs,
        "RAG 工作流程（中文 / 英文）",
        [
            "中文主链路：用户问题 → embedding → 中文医学知识库向量检索 → Top-K 证据拼接 → LLM 生成回答（可附引用与阈值控制）。",
            "英文扩展：独立或并行英文索引，用于补充指南、术语与双语场景，检索流程与中文同构。",
            "质量策略：相似度阈值过滤、无命中时降级为纯上下文回答；支持仅检索模式与调试轨迹。",
            "数据层：文档切块、向量化、持久化存储；可重建索引以接入增量语料。",
        ],
    )

    add_bullet_slide(
        prs,
        "运动识别与三餐识别流程",
        [
            "运动识别：采集图像/视频帧 → 姿态或动作模型推理 → 动作分类与计次/计时 → 会话统计反馈至小程序与档案侧。",
            "三餐识别：食物图像输入 → 营养识别模型 → 热量与宏量营养素等结构化结果 → 结合健康建议形成饮食干预线索。",
            "工程形态：独立训练/推理服务，与主后端 API 对接；便于替换权重与版本迭代。",
            "业务价值：康复训练量化 + 饮食行为闭环，与 RAG、档案形成互补。",
        ],
    )

    add_bullet_slide(
        prs,
        "基于 Heap 的 NCSS 优化",
        [
            "NCSS 要点：用户相似图 → 共同邻居 → 类 PageRank 中心度 → 种子筛选 → 社团扩张（结构强度 S、依赖度 D）→ 小社团合并。",
            "Heap 优化逻辑：扩张阶段在边界候选集中维护「结构强度」大顶堆（或等价优先队列），O(log n) 取当前最优节点，替代反复全量排序。",
            "工程效果：候选规模随网络增大时，单轮扩张的比较与排序开销显著降低，迭代更稳定、更易实时化。",
            "优化表现（汇报口径）：时间复杂度由轮次×候选排序主导降为轮次×堆操作；大规模节点下墙钟时间与 CPU 占用更优（可附对比曲线）。",
            "说明：混合图（疾病集合 Jaccard + 预测向量相似）在 run_ncss_json 中与后端社团服务对齐。",
        ],
    )

    add_bullet_slide(
        prs,
        "向量相似度 · 社团划分 · 资讯推荐 · 互助",
        [
            "输入：疾病预测/分析产生的用户向量与标签，与历史行为特征共同构图。",
            "相似度分析：向量余弦等与图边权融合，供 NCSS 划分稳定社团。",
            "社团运营：同社团用户画像相近，便于精准健康干预与病友匹配。",
            "头条/热点：爬虫脚本抓取与用户兴趣、疾病标签相关的资讯，入库并推荐，增强粘性。",
            "互助闭环：产品侧鼓励同社团病情交流、经验分享，与问答、档案、运动/饮食数据形成长期留存。",
        ],
    )

    base = Path(__file__).resolve().parent
    out = base / "医视智行-项目答辩-8页.pptx"
    prs.save(str(out))
    alt = base / "yishizhixing-pitch-8slides.pptx"
    prs.save(str(alt))
    print("已生成:", out)
    print("副本(英文文件名):", alt)


if __name__ == "__main__":
    main()
